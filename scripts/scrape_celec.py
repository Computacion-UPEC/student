"""Extrae las convocatorias de empleo de todas las unidades de negocio de CELEC EP,
lee los requisitos (formacion y experiencia) de cada PDF / publicacion y marca
las que aplican para un Ingeniero en Computacion.

Uso:  python scrape_celec.py
Genera un solo archivo: celec_empleos.xlsx (hojas "Para Ing. Computacion" y "Todas").
Requiere: pip install pypdf openpyxl
"""
import hashlib
import html
import io
import json
import os
import re
import tempfile
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime

import pypdf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASE = "https://www.celec.gob.ec"
HOY = date.today()
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "celec_empleos.xlsx")
CACHE = os.path.join(tempfile.gettempdir(), "celec_empleos_cache")
os.makedirs(CACHE, exist_ok=True)

# Paginas con tabla (Cargo | Fecha publicacion | Fecha finalizacion | Codigo | Documento)
TABLAS = {
    "CELEC SUR": "/celecsur/bolsa-de-empleo/",
    "Coca Codo Sinclair": "/cocacodo/bolsa-de-empleo/",
    "Hidrotoapi": "/hidrotoapi/bolsa-de-empleo/",
    "Termoesmeraldas": "/termoesmeraldas/oportunides-de-empleo/",
    "Termomanabí": "/termomanabi/bolsa-de-empleo/",
}
# Paginas tipo lista "dd/mm/aaaa – CONVOCATORIA – CARGO" + PDF
LISTAS = {"Termogas Machala": "/termogasmachala/oportunidades-de-empleo/"}
# Todas las unidades: convocatorias publicadas como entradas (posts) de WordPress
POSTS = {
    "CELEC EP Matriz": "", "CELEC SUR": "celecsur", "Coca Codo Sinclair": "cocacodo",
    "Electroguayas": "electroguayas", "Hidroagoyán": "hidroagoyan", "Gensur": "gensur",
    "Hidronación": "hidronacion", "Hidrotoapi": "hidrotoapi", "Termoesmeraldas": "termoesmeraldas",
    "Termogas Machala": "termogasmachala", "Termomanabí": "termomanabi",
    "Termopichincha": "termopichincha", "Transelectric": "transelectric",
}
BUSQUEDAS = ("convocatoria", "selección", "vinculación", "talento humano")
EMPLEO = re.compile(r"convocatoria|selecci[oó]n|vinculaci[oó]n|oferta de trabajo|talento humano", re.I)
NO_EMPLEO = re.compile(r"ofertas|licitaci|feria inclusiva|interconexi|estructurador|concesi|manifestaci|"
                       r"expresiones de inter|contrataci[oó]n del servicio|proveedor|subasta|cotizaci|"
                       r"capacitaci|socializaci[oó]n del|audiencia|rendici[oó]n|"
                       r"invitaci[oó]n a convocatoria p[uú]blica|giro espec[ií]fico|arrendamiento", re.I)

# Convocatorias publicadas solo como imagen (sin texto): requisitos transcritos a mano.
# Clave: enlace de la publicacion.
MANUAL = {
    f"{BASE}/electroguayas/noticias/convocatoria-procesos-de-seleccion-simple-electroguayas-2/": {
        "Cargo": "ASISTENTE ADMINISTRATIVO / ASISTENTE ADM. 7 (Central Aníbal Santos y Álvaro Tinajero, Guayaquil)",
        "fin": date(2026, 10, 1), "Vacantes": "1", "Sueldo (RMU)": "$ 1.170,00",
        "Formación requerida": "Tercer Nivel Tecnológico: Tecnológico Superior ó más de 2 años en título de Tercer "
            "Nivel de Grado. Incluye: Ciencias computacionales; Diseño y administración de redes y bases de datos; "
            "Desarrollo y análisis de software y aplicaciones; Sistemas de información; Innovación tecnológica en "
            "los negocios; (y Administración, Contabilidad, Economía, Derecho, etc.) o carreras afines. "
            "Postular hasta 16h30 del 01/10/2026 a seleccion.egu@celec.gob.ec (ASUNTO: cargo).",
        "Experiencia requerida": "2 años en cargos similares; ó 3 años en cargos de Ayudantes/Auxiliares, "
            "relacionadas con el área.",
    },
    f"{BASE}/electroguayas/noticias/convocatoria-procesos-de-seleccion-simple-electroguayas/": {
        "Cargo": "ESPECIALISTA JURÍDICO / ESPECIALISTA ADM. 4 (Central Gonzalo Zevallos, Guayaquil)",
        "fin": date(2026, 10, 2), "Vacantes": "1", "Sueldo (RMU)": "$ 2.025,00",
        "Formación requerida": "Tercer Nivel de Grado: Licenciatura y títulos profesionales; Mínimo 8 semestres, "
            "Jurisprudencia; o carreras afines en función de la formación.",
        "Experiencia requerida": "3 años en cargos similares, o 4 años en cargos de Asistentes, relacionados con el área.",
    },
}

MESES = {m: i for i, m in enumerate(
    "enero febrero marzo abril mayo junio julio agosto septiembre octubre noviembre diciembre".split(), 1)}
NUM = {"un": 1, "uno": 1, "una": 1, "dos": 2, "tres": 3, "cuatro": 4, "cinco": 5, "seis": 6,
       "siete": 7, "ocho": 8, "nueve": 9, "diez": 10, "doce": 12, "dieciocho": 18, "veinticuatro": 24}

# ---- Elegibilidad para Ingenieria en Computacion -------------------------------------------
DIRECTO = re.compile(r"computaci|inform[aá]tic|software|sistemas(?!\s+el[eé]ctric|\s+de\s+potencia)|"
                     r"tecnolog[ií]as?\s+de\s+(la\s+)?informaci|\bTICs?\b|telem[aá]tic|ciberseguridad|"
                     r"ciencias?\s+de\s+datos|redes\s+y\s+telecom", re.I)
AFIN = re.compile(r"electr[oó]nic|telecomunicaci|mecatr[oó]nic|automatizaci|ingenier[ií]a\s+en\s+cualquier|"
                  r"cualquier\s+(carrera|[aá]rea|disciplina|profesi[oó]n)|ingenier[ií]as?\s*[;,.]", re.I)
BACHILLER = re.compile(r"bachiller", re.I)
TITULO_TIC = re.compile(r"\bTIC\b|tecnolog[ií]a|sistemas|inform[aá]tic|software|redes|datos|"
                        r"automatizaci|scada|telecomunic|electr[oó]nic|instrumentaci|ciberseg|desarrollador", re.I)


def get(url, binario=False):
    clave = os.path.join(CACHE, hashlib.md5(url.encode()).hexdigest())
    if binario and os.path.exists(clave):
        return open(clave, "rb").read()
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=90) as r:
        data = r.read()
    if binario:
        open(clave, "wb").write(data)
        return data
    return data.decode("utf-8", errors="ignore")


def texto(fragmento):
    s = re.sub(r"<br\s*/?>|</p>|</li>", "\n", fragmento)
    s = html.unescape(re.sub(r"<[^>]+>", " ", s))
    return re.sub(r"[ \t\xa0]+", " ", s).strip()


def plano(s):
    return re.sub(r"\s+", " ", s).strip()


def parse_fecha(s):
    if not s:
        return None
    s = s.lower().strip()
    m = re.search(r"(\d{1,2})\s*de\s*([a-záéíóú]+)\s*(?:de|del)?\s*(\d{4})", s)
    if m and m.group(2) in MESES:
        try:
            return date(int(m.group(3)), MESES[m.group(2)], int(m.group(1)))
        except ValueError:
            return None
    m = re.search(r"(\d{1,2})[-/](\d{1,2})[-/](\d{4})", s)
    if m:
        try:
            return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    return None


def fila(unidad, cargo, pub=None, fin=None, codigo="", vacantes="", doc="", fuente="", contenido=""):
    return {"Unidad": unidad, "Cargo": plano(cargo), "Vacantes": vacantes, "Fecha publicación": pub,
            "Fecha finalización": fin, "Código": codigo, "Documento": doc, "Fuente": fuente,
            "_contenido": contenido}


# ---- Extraccion de convocatorias ----------------------------------------------------------
def scrape_tablas():
    out = []
    for unidad, path in TABLAS.items():
        url = BASE + path
        s = get(url)
        for tabla in re.findall(r"<table.*?</table>", s, re.S):
            filas = re.findall(r"<tr.*?</tr>", tabla, re.S)
            if not filas:
                continue
            cab = [plano(texto(c)).lower() for c in re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", filas[0], re.S)]

            def col(*claves):
                return next((i for i, c in enumerate(cab) if any(k in c for k in claves)), None)

            i_cargo, i_pub, i_fin = col("cargo"), col("publica"), col("finaliza")
            i_cod, i_vac = col("código", "codigo"), col("vacante")
            if i_cargo is None:
                continue
            for tr in filas[1:]:
                v = [plano(texto(c)) for c in re.findall(r"<td[^>]*>(.*?)</td>", tr, re.S)]
                if len(v) <= i_cargo or not v[i_cargo] or v[i_cargo].lower() == "cargo":
                    continue
                g = lambda i: v[i] if i is not None and i < len(v) else ""
                pdf = re.search(r'href="([^"]+\.pdf)"', tr, re.I) or re.search(r'href="([^"]+)"', tr)
                out.append(fila(unidad, v[i_cargo], parse_fecha(g(i_pub)), parse_fecha(g(i_fin)),
                                g(i_cod), g(i_vac), pdf.group(1) if pdf else "", url))
    return out


def scrape_listas():
    out = []
    sep = r"\s*(?:–|-|&#8211;)\s*"
    for unidad, path in LISTAS.items():
        url = BASE + path
        s = get(url)
        patron = re.compile(r"(\d{1,2}/\d{1,2}/\d{4})" + sep + "CONVOCATORIA" + sep + r"(.*?)<(.*?)"
                            r"(?=\d{1,2}/\d{1,2}/\d{4}" + sep + "CONVOCATORIA|$)", re.S)
        for f, cargo, resto in patron.findall(s):
            pdf = re.search(r'href="([^"]+\.pdf)"', resto, re.I)
            out.append(fila(unidad, texto(cargo), parse_fecha(f), doc=pdf.group(1) if pdf else "", fuente=url))
    return out


def scrape_matriz():
    url = BASE + "/bolsa-de-empleo/"
    s = get(url)
    out = []
    for cargo, pdf, etiqueta in re.findall(
            r"<(?:h\d|p)[^>]*>((?:(?!</?(?:h\d|p)\b).)*?)</(?:h\d|p)>\s*<div class=\"wp-block-file\">"
            r"<a[^>]+href=\"([^\"]+\.pdf)\"[^>]*>(.*?)</a>", s, re.S):
        cargo, etiqueta = plano(texto(cargo)), plano(texto(etiqueta))
        if cargo and "convocatoria" in etiqueta.lower():
            out.append(fila("CELEC EP Matriz", cargo, parse_fecha(etiqueta), doc=pdf, fuente=url))
    return out


def posts_unidad(unidad, slug):
    raiz = f"{BASE}/{slug}".rstrip("/")
    vistos, out = set(), []
    for termino in BUSQUEDAS:
        pagina = 1
        while True:
            api = (f"{raiz}/wp-json/wp/v2/posts?search={urllib.parse.quote(termino)}"
                   f"&per_page=100&page={pagina}&_fields=id,date,link,title,content")
            try:
                posts = json.loads(get(api))
            except Exception:
                break
            if not isinstance(posts, list) or not posts:
                break
            for p in posts:
                titulo = plano(texto(p["title"]["rendered"]))
                if p["id"] in vistos or not EMPLEO.search(titulo) or NO_EMPLEO.search(titulo):
                    continue
                vistos.add(p["id"])
                # "... invita a participar en el proceso de seleccion simple para ocupar el puesto de X"
                m = re.search(r"(?:puesto|cargo)s?\s+de\s*:?\s*(.+)", titulo, re.I)
                if m and re.match(r"CELEC EP \w+ invita", titulo, re.I):
                    titulo = m.group(1).strip(" .")
                out.append(fila(unidad, titulo, datetime.fromisoformat(p["date"]).date(), doc=p["link"],
                                fuente=f"{raiz}/ (publicaciones)", contenido=p["content"]["rendered"]))
            if len(posts) < 100:
                break
            pagina += 1
    return out


def scrape_posts():
    with ThreadPoolExecutor(6) as ex:
        return [f for r in ex.map(lambda kv: posts_unidad(*kv), POSTS.items()) for f in r]


# ---- Requisitos ----------------------------------------------------------------------------
def texto_pdf(url):
    try:
        r = pypdf.PdfReader(io.BytesIO(get(url, binario=True)))
        return "\n".join((p.extract_text() or "") for p in r.pages)
    except Exception:
        return ""


def texto_documento(f):
    """Texto completo con los requisitos de la convocatoria (PDF o publicacion + sus PDFs)."""
    if f["_contenido"]:
        c = f["_contenido"]
        t = texto(c)
        for pdf in re.findall(r'href="([^"]+\.pdf)"', c, re.I)[:6]:
            t += "\n" + texto_pdf(pdf)
        return t
    if f["Documento"].lower().endswith(".pdf"):
        return texto_pdf(f["Documento"])
    return ""


INI_FORM = re.compile(r"(formaci[oó]n acad[eé]mica|instrucci[oó]n formal\s*:|tercer nivel|cuarto nivel|"
                      r"bachiller|t[eé]cnico superior|tecnol[oó]gico superior|t[ií]tulo de)", re.I)
INI_EXP = re.compile(r"(tiempo de experiencia|experiencia\s*:|alternativa\s*1|no requiere experiencia|"
                     r"sin experiencia|\d+\s*(a[nñ]os?|meses)\s+(en|de)\s+(cargos|experiencia))", re.I)


def meses_experiencia(s):
    """Minimo de meses entre las alternativas citadas; 0 si no requiere."""
    s = s.lower()
    if re.search(r"no requiere|sin experiencia|no se requiere|no aplica", s):
        return 0
    valores = []
    for n, unidad in re.findall(r"(\d+(?:[.,]\d+)?|" + "|".join(NUM) + r")\s*\(?\d*\)?\s*(a[nñ]os?|meses?)", s):
        n = NUM.get(n) or float(n.replace(",", "."))
        valores.append(round(n * 12) if unidad.startswith("a") else round(n))
    return min(valores) if valores else None


def requisitos(f, t):
    """Devuelve (formacion, experiencia_texto, meses, sueldo) para el cargo de la fila."""
    t = t.replace("\r", "")
    plano_t = plano(t)
    # Si la publicacion trae varios cargos ("Cargo: X ... Cargo: Y"), quedarse con el bloque del cargo
    bloques = re.split(r"(?=\bCargo\s*:)", plano_t)
    if len(bloques) > 2:
        palabras = [w for w in re.findall(r"\w{5,}", f["Cargo"].lower()) if w not in ("convocatoria", "proceso",
                    "selecci", "simple", "vinculaci", "central")]
        mejor = max(bloques[1:], key=lambda b: sum(w in b.lower() for w in palabras))
        bloques_cargo = [mejor] if palabras else bloques[1:]
    else:
        bloques_cargo = [plano_t]
    # Si es una publicacion generica de varios cargos, evaluamos todos sus bloques
    if len(bloques) > 2 and re.search(r"varios|procesos de selecci[oó]n simple\s*(electroguayas)?$|"
                                      r"invitaci", f["Cargo"], re.I):
        bloques_cargo = bloques[1:]

    forms, exps, meses, sueldos = [], [], [], []
    for b in bloques_cargo:
        mf = INI_FORM.search(b)
        me = INI_EXP.search(b, mf.end() if mf else 0)
        if mf:
            fin_f = me.start() if me and me.start() > mf.start() else mf.start() + 350
            forms.append(b[mf.start():fin_f][:450].strip(" -–:"))
        if me:
            seg = b[me.start():me.start() + 260]
            seg = re.split(r"Versi[oó]n|M15\.|Los aspirantes|Cargo\s*:|PROCESOS DE", seg)[0]
            exps.append(seg.strip(" -–:"))
            m = meses_experiencia(seg)
            if m is not None:
                meses.append(m)
        sm = re.search(r"\$\s?([\d.]+,\d{2}|\d[\d.,]*)", b) or re.search(r"\b(\d{3,4}\.\d{2})\b", b)
        if sm:
            sueldos.append("$ " + sm.group(1))
    return (" || ".join(dict.fromkeys(forms)), " || ".join(dict.fromkeys(exps)),
            min(meses) if meses else None, ", ".join(dict.fromkeys(sueldos[:3])))


def elegibilidad(cargo, formacion):
    if DIRECTO.search(formacion):
        return "1. Directa (Computación/Sistemas/TI)"
    if AFIN.search(formacion) or (TITULO_TIC.search(cargo) and "afines" in formacion.lower()):
        return "2. Posible (carrera afín / ingeniería)"
    if not formacion and TITULO_TIC.search(cargo):
        return "2. Posible (cargo TIC, revisar requisitos)"
    if BACHILLER.search(formacion) and not re.search(r"tercer nivel|tecnol[oó]gico", formacion, re.I):
        return "3. Solo bachiller (cualquiera puede aplicar)"
    if not formacion:
        return "Sin requisitos legibles - revisar"
    return "No aplica"


# ---- Main ----------------------------------------------------------------------------------
def main():
    filas = []
    for nombre, fn in [("tablas", scrape_tablas), ("listas", scrape_listas),
                       ("matriz", scrape_matriz), ("publicaciones", scrape_posts)]:
        try:
            r = fn()
            print(f"{nombre}: {len(r)}")
            filas += r
        except Exception as e:
            print(f"{nombre}: ERROR {e}")

    vistos, unicas = set(), []
    for f in filas:
        k = (f["Unidad"], f["Cargo"].lower(), f["Documento"])
        if k not in vistos:
            vistos.add(k)
            unicas.append(f)
    print(f"Leyendo requisitos de {len(unicas)} convocatorias...")

    with ThreadPoolExecutor(8) as ex:
        textos = list(ex.map(texto_documento, unicas))
    for f, t in zip(unicas, textos):
        form, exp, meses, sueldo = requisitos(f, t)
        if f["Documento"] in MANUAL:
            m = MANUAL[f["Documento"]]
            f["Cargo"], f["Fecha finalización"], f["Vacantes"] = m["Cargo"], m["fin"], m["Vacantes"]
            form, exp, sueldo = m["Formación requerida"], m["Experiencia requerida"], m["Sueldo (RMU)"]
            meses = meses_experiencia(exp)
        elif not form and re.search(r"<img", f["_contenido"]):
            exp = "Requisitos publicados como imagen: abrir el enlace"
        f.update({"Formación requerida": form, "Experiencia requerida": exp,
                  "Experiencia mínima (meses)": meses, "Sueldo (RMU)": sueldo,
                  "Elegibilidad Ing. Computación": elegibilidad(f["Cargo"], form)})
        fin, pub = f["Fecha finalización"], f["Fecha publicación"]
        if fin:
            f["Estado"] = "Abierta" if fin >= HOY else "Cerrada"
        elif pub and (HOY - pub).days <= 7:
            f["Estado"] = "Reciente - revisar"
        else:
            f["Estado"] = "Cerrada (sin fecha de cierre)"

    cols = ["Unidad", "Cargo", "Elegibilidad Ing. Computación", "Experiencia mínima (meses)",
            "Experiencia requerida", "Formación requerida", "Sueldo (RMU)", "Vacantes", "Fecha publicación",
            "Fecha finalización", "Estado", "Código", "Documento"]
    anchos = [18, 50, 30, 13, 55, 70, 14, 9, 13, 13, 16, 18, 9]

    def orden_exp(f):
        m = f["Experiencia mínima (meses)"]
        return (m if m is not None else 999, -(f["Fecha publicación"] or date.min).toordinal())

    aplica = sorted([f for f in unicas if f["Elegibilidad Ing. Computación"].startswith(("1", "2"))], key=orden_exp)
    todas = sorted(unicas, key=lambda f: f["Fecha publicación"] or date.min, reverse=True)

    wb = Workbook()
    hojas = [("Para Ing. Computación", aplica), ("Todas", todas)]
    verde, amarillo = PatternFill("solid", fgColor="C6EFCE"), PatternFill("solid", fgColor="FFF2CC")
    for i, (nombre, datos) in enumerate(hojas):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = nombre
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for f in datos:
            ws.append([f[c] for c in cols])
            r = ws.max_row
            for c in (9, 10):
                ws.cell(r, c).number_format = "DD/MM/YYYY"
            for c in (5, 6):
                ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
            if f["Documento"]:
                ws.cell(r, 13).hyperlink = f["Documento"]
                ws.cell(r, 13).value = "Abrir"
                ws.cell(r, 13).font = Font(color="0563C1", underline="single")
            if f["Estado"] in ("Abierta", "Reciente - revisar"):
                for c in range(1, len(cols) + 1):
                    ws.cell(r, c).fill = verde
            elif f["Experiencia mínima (meses)"] == 0:
                ws.cell(r, 4).fill = amarillo
        for idx, ancho in enumerate(anchos, 1):
            ws.column_dimensions[ws.cell(1, idx).column_letter].width = ancho
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(OUT)

    from collections import Counter
    print(f"\nTotal: {len(unicas)} | por unidad: {dict(Counter(f['Unidad'] for f in unicas))}")
    print(f"Elegibilidad: {dict(Counter(f['Elegibilidad Ing. Computación'] for f in unicas))}")
    print(f"Para Ing. Computación: {len(aplica)} | archivo: {OUT}")


if __name__ == "__main__":
    main()
